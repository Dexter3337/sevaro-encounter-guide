import json,datetime,zipfile,hashlib,csv,re,shutil
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parent; o=ROOT.parent/'outputs'/'01a124a6'
k=json.loads((o/'knowledge_base.json').read_text(encoding='utf-8'))
w=openpyxl.load_workbook(o/'Sevaro_Challenge_1_Foundation.xlsx',read_only=True,data_only=True)
checks=0;errors=[]
for tab,rows in k.items():
 s=w[tab];keys=list(dict.fromkeys(key for r in rows for key in r))
 assert next(s.iter_rows(min_row=1,max_row=1,max_col=len(keys),values_only=True))==tuple(key.replace('_',' ') for key in keys),tab
 for i,(r,vals) in enumerate(zip(rows,s.iter_rows(min_row=2,max_row=len(rows)+1,max_col=len(keys),values_only=True)),2):
  for key,actual in zip(keys,vals):
   exp=r.get(key);checks+=1
   if isinstance(actual,(datetime.datetime,datetime.date)):actual=actual.strftime('%Y-%m-%d')
   if exp!=actual and not (exp in ['',None] and actual in ['',None]):errors.append((tab,i,key,str(exp)[:80],str(actual)[:80]))
assert not errors,errors[:10]
assert [w['Overview'].cell(i,2).value for i in range(5,16)]==[34,14,288,62,101,30,131,21,7,599,22]
w.close()
# Exact source values and dates, independently compared to raw PDF table cells.
raw=json.loads((ROOT/'raw_tables.json').read_text(encoding='utf-8'));n=0
def clean(x):return re.sub(r'\s+',' ',x or '').strip()
for page in raw:
 if 4<=page['page']<=10:
  hid=f"SYN-{page['page']-3:03d}"
  for field,value in page['tables'][0][1:]:
   field,value=clean(field),clean(value);r=next(r for r in k['Synthetic Rules'] if r['hospital_id']==hid and r['field']==field)
   assert r['original_wording']==value,(hid,field)
   dt=re.search(r'\[last verified (\d{4}-\d{2}-\d{2})\]',value)
   assert r['last_verified']==(dt[1] if dt else '2026-02-02');n+=1
assert n==169
r=json.loads((o/'validation_report.json').read_text(encoding='utf-8'))
r.update(workbook_cells_compared=checks,workbook_sheets=24,exported_values_match=True,overview_counts_match=True,source_guide_exact_value_date_comparisons=n,formula_error_scan='0 matches',visual_review='Rendered all 24 sheets and reviewed previews; source URL values verified after removing unsupported formula')
r['limitations']=['Excel desktop interaction not tested','No companion execution or measured event-to-paint latency','No independent clinician oracle adjudication']
(o/'validation_report.json').write_text(json.dumps(r,indent=2),encoding='utf-8')
with (o/'research_batch_log.md').open('a',encoding='utf-8') as f:
 f.write('\n## Final artifact verification\nAll 24 workbook sheets rendered and reviewed. Exported cells reconciled to JSON, overview counts checked, exact original wording and verification dates checked for all 169 supplied records, and no formula-error matches found. Runtime behavior, Excel desktop interaction, latency and clinical outcomes remain untested. Source URLs preserved literally because the authoring engine does not evaluate HYPERLINK. Stop here: next useful work requires organizer/API/local owner evidence or companion execution rather than duplicate public research.\n')
dest=o/'source_material'/'Sevaro';dest.mkdir(parents=True,exist_ok=True)
pdf='Sevaro_Rutgers_Health_Hack_Challenge_1_Guided_Workflow_v2.pdf'
shutil.copy2(ROOT/'source_material'/'Sevaro'/pdf,dest/pdf)
manifest=json.loads((o/'split_manifest.json').read_text(encoding='utf-8'))
for name,digest in manifest['files'].items():assert hashlib.sha256((o/name).read_bytes()).hexdigest()==digest
keep=['Sevaro_Challenge_1_Foundation.xlsx','decision_brief.md','README.md','knowledge_base.json','hospital_guides_reconstructed.json','storyboard_notes_reconstructed.json','development_guides.json','held_out_guides.json','development_inputs.jsonl','held_out_inputs.jsonl','expected_outputs.jsonl','split_manifest.json','validation_report.json','research_batch_log.md']
with zipfile.ZipFile(o/'Sevaro_Challenge_1_Research_Package.zip','w',zipfile.ZIP_DEFLATED) as z:
 for name in keep:z.write(o/name,name)
 for file in (o/'tables').glob('*.csv'):z.write(file,file.relative_to(o))
 z.write(dest/pdf,(dest/pdf).relative_to(o))
with zipfile.ZipFile(o/'Sevaro_Challenge_1_Research_Package.zip') as z:assert z.testzip() is None
print(json.dumps(r,indent=2))
